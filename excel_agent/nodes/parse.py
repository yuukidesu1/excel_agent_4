"""
nodes/parse.py — Excel 结构解析节点（纯代码）

职责：
    1. 读取目标 Sheet 的所有非空单元格（含位置、值、粗体标记）
    2. 收集全部合并单元格信息（范围、值、跨行/列数）
    3. 基于"连续空行"将 Sheet 粗粒度切分为候选子表块
       → 这正是灵犀"初步拆分出几个表格"的实现方式

输出放入 state["sheet_structure"]，供 locate_node 使用。
不调用 LLM，执行速度快，且只执行一次（重试时跳过）。
"""

import openpyxl
from typing import List, Dict, Any
from Excel_Agent.excel_agent.state import AgentState


def _find_sheet(wb, requested: str) -> str:
    """大小写不敏感 + 去首尾空格匹配 Sheet 名"""
    if requested in wb.sheetnames:
        return requested
    req = requested.strip().upper()
    for name in wb.sheetnames:
        if name.strip().upper() == req:
            return name
    return requested  # 找不到时返回原值，让 openpyxl 抛出明确错误


def _row_is_empty(ws, row_idx: int, max_col: int) -> bool:
    for c in range(1, max_col + 1):
        if ws.cell(row=row_idx, column=c).value is not None:
            return False
    return True


def parse_node(state: AgentState) -> dict:
    # try:
    #     wb         = openpyxl.load_workbook(state["config"]["excel_path"])
    # except Exception as e:
    #     return {"error" : f"无法读取 Excel 文件: {state['config']['excel_path']}"}
    wb = openpyxl.load_workbook(state["config"]["excel_path"])
    sheet_name = _find_sheet(wb, state["config"]["sheet_name"])
    if sheet_name not in wb.sheetnames:
        return {"error": f"未找到目标 Sheet 页: {state['config']['sheet_name']}"}
    ws         = wb[sheet_name]
    max_row    = ws.max_row
    max_col    = ws.max_column

    # ── 1. 遍历所有非空单元格 ────────────────────────────────
    non_empty_cells:   List[Dict] = []
    potential_headers: List[Dict] = []
    merged_coords = set()

    # 获取用户配置
    usr_config = state.get("config", "")
    # 获取用户配置的表头
    usr_titles = usr_config.get("subtable_titles")
    # 先收集合并单元格坐标
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                merged_coords.add((r, c))

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            is_bold   = bool(cell.font and cell.font.bold)
            is_merged = (cell.row, cell.column) in merged_coords
            info = {
                "coord":  cell.coordinate,
                "row":    cell.row,
                "col":    cell.column,
                "value":  str(cell.value).replace('\n', ' ').replace('\r', '').strip()[:300],    # 清洗换行符
                "bold":   is_bold,
                "merged": is_merged,
            }
            non_empty_cells.append(info)
            # 粗体 or 合并单元格 → 候选表头
            if is_bold or is_merged:
                potential_headers.append(info)

    # ── 2. 合并单元格详情 ────────────────────────────────────
    merged_cells_info: List[Dict] = []
    for m in ws.merged_cells.ranges:
        top_val = ws.cell(row=m.min_row, column=m.min_col).value
        merged_cells_info.append({
            "range":    str(m),
            "min_row":  m.min_row,
            "max_row":  m.max_row,
            "min_col":  m.min_col,
            "max_col":  m.max_col,
            "value":    str(top_val).replace('\n', ' ').replace('\r', '').strip() if top_val is not None else "",
            "row_span": m.max_row - m.min_row + 1,
            "col_span": m.max_col - m.min_col + 1,
        })

    # ── 3. 基于空行切分候选子表块 ────────────────────────────
    # 与灵犀"根据空间分布初步拆分"逻辑一致
    potential_subtables: List[Dict] = []
    in_block    = False
    block_start = 1

    for r in range(1, max_row + 2):
        is_empty = (r > max_row) or _row_is_empty(ws, r, max_col)

        if not is_empty and not in_block:
            block_start = r
            in_block    = True
        elif is_empty and in_block:
            title_candidates = [
                c["value"] for c in non_empty_cells
                if block_start <= c["row"] <= block_start + 2
            ]
            for usr_title in (usr_titles or []):
                title_candidates.append(usr_title)
            potential_subtables.append({
                "start_row":        block_start,
                "end_row":          r - 1,
                "row_count":        r - 1 - block_start + 1,
                "title":            title_candidates[0] if title_candidates else "",
                "title_candidates": title_candidates[:15],
            })
            in_block = False

    # DEBUG
    res = {
        "sheet_structure": {
            "sheet_name": sheet_name,
            "max_row": max_row,
            "max_col": max_col,
            "non_empty_cells": non_empty_cells,
            "potential_headers": potential_headers,
            "merged_cells_info": merged_cells_info,
            "potential_subtables": potential_subtables,
        }
    }
    return res
    # return {
    #     "sheet_structure": {
    #         "sheet_name":          sheet_name,
    #         "max_row":             max_row,
    #         "max_col":             max_col,
    #         "non_empty_cells":     non_empty_cells,
    #         "potential_headers":   potential_headers,
    #         "merged_cells_info":   merged_cells_info,
    #         "potential_subtables": potential_subtables,
    #     }
    # }