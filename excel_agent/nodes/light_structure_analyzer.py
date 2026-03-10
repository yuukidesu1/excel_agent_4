"""
nodes/light_structure_analyzer.py — 轻量级结构分析节点

职责：
    在 parse 之后立即执行，基于 parse 输出计算一个轻量级结构指纹，
    用于 L2 缓存查询（在 code_gen 之前）。

与 structure_analyzer_node 的区别：
    - 不调用 LLM，纯代码计算
    - 基于 parse 输出的 sheet_structure 直接计算
    - 精度较低，但速度极快（~5ms）
    - 用于 L2 缓存的早期查询
"""

import hashlib
import json
from typing import Dict, List, Any, Optional

from excel_agent.state import AgentState


def _normalize_title(title: str) -> str:
    """归一化标题：小写、去空格、去标点"""
    import re
    s = title.lower()
    s = re.sub(r'\s+', '', s)
    s = re.sub(r'[^\w\u4e00-\u9fff]', '', s)
    return s


def _compute_header_signature(ws, start_row: int, header_rows: int, max_col: int) -> str:
    """
    计算表头签名（基于表头内容）

    读取 header_rows 行的内容，计算哈希值
    用于区分结构相同但表头不同的情况
    """
    try:
        header_content = []
        for r in range(start_row, start_row + header_rows):
            row_data = []

            # [FIX 1] 计算相对行号，确保表头整体上下平移时，哈希值保持不变
            rel_r = r - start_row

            for c in range(1, min(max_col + 1, 30)):  # 限制列数防止过大
                cell = ws.cell(row=r, column=c)
                val = cell.value
                if val is not None:
                    # 归一化：转小写、去空格
                    normalized = str(val).lower().strip()
                    # 使用相对行号 rel_r 进行字符串拼接
                    row_data.append(f"R{rel_r}C{c}:{normalized}")
            header_content.append("|".join(row_data))

        content_str = "|||".join(header_content)
        return hashlib.sha256(content_str.encode()).hexdigest()[:12]
    except Exception:
        return "unknown"


def _detect_header_rows(ws, start_row: int, max_col: int) -> int:
    """
    检测表头行数
    """
    # 检查第一行
    row1_has_header = False
    row2_has_header = False

    try:
        # 检查 start_row 行
        for c in range(1, min(max_col + 1, 20)):
            cell1 = ws.cell(row=start_row, column=c)
            if cell1.font and cell1.font.bold:
                row1_has_header = True
                break
            # 检查是否在合并单元格范围内
            for merged_range in ws.merged_cells.ranges:
                if merged_range.min_row <= start_row <= merged_range.max_row:
                    row1_has_header = True
                    break

        # 检查 start_row+1 行
        for c in range(1, min(max_col + 1, 20)):
            cell2 = ws.cell(row=start_row + 1, column=c)
            if cell2.font and cell2.font.bold:
                row2_has_header = True
                break
            for merged_range in ws.merged_cells.ranges:
                if merged_range.min_row <= (start_row + 1) <= merged_range.max_row:
                    row2_has_header = True
                    break

        # 检查第二行是否有表头关键词
        for c in range(1, min(max_col + 1, 20)):
            cell2 = ws.cell(row=start_row + 1, column=c)
            val = str(cell2.value).lower() if cell2.value else ""
            if any(kw in val for kw in ['type', 'qty', 'name', 'id', 'no', 'module', 'cell']):
                row2_has_header = True
                break
    except Exception:
        pass

    if row1_has_header and row2_has_header:
        return 2
    elif row1_has_header:
        return 1
    else:
        return 1  # 默认 1 行表头


def light_structure_analyzer_node(state: AgentState) -> dict:
    """
    轻量级结构分析节点
    """
    import openpyxl

    config = state.get("config", {})
    excel_path = config.get("excel_path")
    sheet_name = config.get("sheet_name")
    subtable_titles = config.get("subtable_titles", [])

    st = state.get("sheet_structure", {})

    if not excel_path or not st:
        return {
            "light_structure_fingerprint": None,
            "light_analyzer_error": "缺少 excel_path 或 sheet_structure"
        }

    try:
        wb = openpyxl.load_workbook(excel_path)
        ws = wb[sheet_name]
    except Exception as e:
        return {
            "light_structure_fingerprint": None,
            "light_analyzer_error": f"无法加载 Excel: {e}"
        }

    max_col = st.get("max_col", ws.max_column)
    potential_subtables = st.get("potential_subtables", [])

    fingerprint = {
        "sheet_name": sheet_name,
        "subtables": [],
        "merge_patterns": []
    }

    for target_title in subtable_titles:
        matched_sub = None
        for sub in potential_subtables:
            sub_title = sub.get("title", "")
            title_candidates = sub.get("title_candidates", [])

            norm_target = _normalize_title(target_title)
            if norm_target in _normalize_title(sub_title) or \
                    any(norm_target in _normalize_title(tc) for tc in title_candidates):
                matched_sub = sub
                break

        if not matched_sub:
            continue

        start_row = matched_sub.get("start_row", 0)
        end_row = matched_sub.get("end_row", 0)
        row_count = matched_sub.get("row_count", 0)

        header_rows = _detect_header_rows(ws, start_row, max_col)
        header_signature = _compute_header_signature(ws, start_row, header_rows, max_col)
        data_rows = row_count - header_rows if row_count > header_rows else 0

        col_count = 0
        for c in range(1, max_col + 1):
            has_header_content = False
            # 只扫描表头所在行
            for r in range(start_row, start_row + header_rows):
                cell = ws.cell(row=r, column=c)
                if cell.value is not None and str(cell.value).strip() !="":
                    has_header_content = True
                    break
            if has_header_content:
                col_count += 1

        fingerprint["subtables"].append({
            "title_pattern": _normalize_title(target_title),
            "original_title": target_title,
            "start_row": start_row,
            "header_rows": header_rows,
            "data_rows": data_rows,
            "row_count": row_count,
            "col_count": col_count,
            "header_signature": header_signature,
        })

    # [FIX 3] 提取并排序合并单元格范围，确保顺序的确定性，避免哈希结果的随机性
    sorted_ranges = sorted(
        ws.merged_cells.ranges,
        key=lambda x: (x.min_row, x.min_col)
    )

    merge_patterns = []
    for merged_range in sorted_ranges:
        row_span = merged_range.max_row - merged_range.min_row + 1
        col_span = merged_range.max_col - merged_range.min_col + 1

        if row_span > 1 or col_span > 1:
            if fingerprint["subtables"]:
                ref_row = fingerprint["subtables"][0]["start_row"]
                rel_row = merged_range.min_row - ref_row
                merge_patterns.append({
                    "rel_row": rel_row,
                    "rel_col": merged_range.min_col - 1,
                    "row_span": row_span,
                    "col_span": col_span,
                })

    fingerprint["merge_patterns"] = merge_patterns[:20]

    # [FIX 2] 构建用于计算签名的结构数据时，剔除不应影响结构指纹的 data_rows
    signature_data = {
        "sheet_name": sheet_name,
        "subtables": [
            {
                "title_pattern": sub["title_pattern"],
                "header_rows": sub["header_rows"],
                "col_count": sub["col_count"],
                "header_signature": sub["header_signature"],
            }
            for sub in fingerprint["subtables"]
        ],
        "merge_patterns": fingerprint["merge_patterns"]
    }

    overall_signature = hashlib.sha256(
        json.dumps(signature_data, sort_keys=True).encode()
    ).hexdigest()[:16]

    fingerprint["signature"] = overall_signature

    return {
        "light_structure_fingerprint": fingerprint,
        "light_analyzer_error": None
    }


def compute_light_signature(fingerprint: Dict[str, Any]) -> str:
    """
    从轻量级结构指纹计算签名
    """
    if not fingerprint:
        return ""

    return fingerprint.get("signature", "")