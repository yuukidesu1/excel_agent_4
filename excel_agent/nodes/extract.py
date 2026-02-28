"""
nodes/extract.py — 精确列抽取节点（纯代码）

职责：
    1. 按 header_map.all_columns 的列号顺序，精确读取每列数据
    2. 预建"合并填充图"：合并区域内所有格都映射到左上角的值
       → 避免 openpyxl 读合并区域非左上角格时返回 None
    3. 向前填充（forward-fill）：对 merge_fill_keys 中的列，
       空值时继承上一行的非空值
       → 这正是灵犀处理 SYSTEM MODULE 跨行合并的核心逻辑
    4. 保留所有空单元格（None），restore_node 统一转为 ""
    5. 全部为空时保留一行全 None（Prompt 规则：不能丢失空行）
"""

import openpyxl
from typing import List, Optional, Any, Dict
from excel_agent.state import AgentState


# ─────────────────────────────────────────────────────────────
# 内部工具
# ─────────────────────────────────────────────────────────────

def _build_merge_map(ws) -> Dict[str, Any]:
    """
    构建合并单元格扩展填充映射。

    问题：openpyxl 合并区域内只有左上角格有值，其余格 value = None。
    解决：把左上角的值映射到区域内每一格。

    返回：{ "R{row}C{col}" → 实际值 }
    """
    fill: Dict[str, Any] = {}
    for m in ws.merged_cells.ranges:
        top = ws.cell(row=m.min_row, column=m.min_col).value
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                fill[f"R{r}C{c}"] = top
    return fill


def _read(ws, row: int, col: int, merge_map: Dict[str, Any]) -> Any:
    """读取单元格，自动处理合并单元格"""
    key = f"R{row}C{col}"
    return merge_map.get(key, ws.cell(row=row, column=col).value)


# ─────────────────────────────────────────────────────────────
# 节点函数
# ─────────────────────────────────────────────────────────────

def extract_node(state: AgentState) -> dict:
    """
    输入：header_map + sheet_structure + excel_path
    输出：raw_data（二维列表，不含表头行，空值为 None）
    """
    hmap       = state["header_map"]
    all_cols   = hmap["all_columns"]        # [{"parent", "child", "col_idx"}, ...]
    col_map    = hmap["column_map"]
    fill_keys  = set(hmap["merge_fill_keys"])  # 需要向前填充的列 key

    wb = openpyxl.load_workbook(state["config"]["excel_path"])
    ws = wb[state["sheet_structure"]["sheet_name"]]
    merge_map = _build_merge_map(ws)

    # ── 目标列号列表（按 all_columns 顺序）────────────────────
    col_keys: List[str] = []
    col_indices: List[Optional[int]] = []

    for col in all_cols:
        parent = col.get("parent") or "None"
        key    = f"{parent}||{col['child']}"
        col_keys.append(key)
        # 优先用 col_idx（LLM 直接给出），否则查 column_map
        idx = col.get("col_idx") or col_map.get(key)
        col_indices.append(idx)

    # 需要 forward-fill 的列位置
    fill_positions = {i for i, k in enumerate(col_keys) if k in fill_keys}

    # ── 数据行范围（跳过表头行）──────────────────────────────
    data_start = hmap["subtable_start_row"] + hmap["header_row_count"]
    data_end   = hmap["subtable_end_row"]

    # ── 逐行抽取 + forward-fill ───────────────────────────────
    raw_data: List[List] = []
    prev_vals: List[Any] = [None] * len(col_indices)  # 用于 forward-fill

    for r in range(data_start, data_end + 1):
        row_vals = []
        for i, col_idx in enumerate(col_indices):
            if col_idx is None:
                val = None
            else:
                val = _read(ws, r, col_idx, merge_map)

            # forward-fill：该列标记为需填充 且 当前值为空 → 继承上一行值
            if i in fill_positions and (val is None or str(val).strip() == ""):
                val = prev_vals[i]
            else:
                prev_vals[i] = val  # 更新"上一个非空值"

            row_vals.append(val)
        raw_data.append(row_vals)

    # 全部为空时保留一行
    if not raw_data:
        raw_data = [[None] * len(col_indices)]

    return {"raw_data": raw_data}
